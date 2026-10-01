"""Read the official Census synthesis workbook without summing hierarchy totals."""

from __future__ import annotations

import unicodedata
import zipfile
from collections.abc import Iterator, Sequence
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook

# Explicit mappings from the official 2021 XLSX labels to the existing text contract.
_HEADER_ALIASES = {
    "FREGUESIA": "DTMNFR21",
    "N_EDIFICIOS_CONSTR_ANTES_1945": "N_EDIFICIOS_CONSTR_ANTES 1945",
    "N_EDIFICIOS_COM_NECESSIDADES_REPARACAO": "N_EDIFICIOS_COM NECESSIDADES REPARAÇAO",
    "N_ALOJAMENTOS_FAM_CLASS_RHABITUAL": "N_ALOJAMENTOS_ FAM_CLASS_RHABITUAL",
    "N_ALOJAMENTOS_FAM_CLASS_VAGOS_OU_RESID_SECUNDARIA": (
        "N_ALOJAMENTOS_ FAM_CLASS_VAGOS OU RESID SECUNDARIA"
    ),
    "N_AGREGADOS_DOMESTICOS_PRIVADOS": "N_AGREGADOS DOMESTICOS PRIVADOS",
    "N_INDIVIDUOS_0_14": "N_INDIVIDUOS_0A14",
    "N_INDIVIDUOS_15_24": "N_INDIVIDUOS_15A24",
    "N_INDIVIDUOS_25_64": "N_INDIVIDUOS_25A64",
}


class CensusWorkbookError(ValueError):
    """Raised when a workbook cannot provide unambiguous subsection records."""


def find_workbook_table(
    archive: zipfile.ZipFile,
    required_columns: set[str],
) -> tuple[str, Iterator[dict[str, str | None]]]:
    """Find a synthesis worksheet and stream only its lowest geographic level.

    INE's workbook has a title row, then headers, then national, regional,
    municipal, parish, section, and subsection rows in the same worksheet.
    SUBSECCAO must be present and non-empty to include a row in aggregation.
    """
    for member in archive.infolist():
        if member.is_dir() or Path(member.filename).suffix.casefold() != ".xlsx":
            continue
        workbook = load_workbook(BytesIO(archive.read(member)), read_only=True, data_only=True)
        try:
            for sheet in workbook:
                rows = sheet.iter_rows(values_only=True)
                for row_number, values in enumerate(rows, start=1):
                    if row_number > 10:
                        break
                    headers = tuple(_header(value) for value in values)
                    if not (required_columns | {"DTMNFR21", "SUBSECCAO"}).issubset(headers):
                        continue
                    nonempty = [header for header in headers if header]
                    if len(set(nonempty)) != len(nonempty):
                        raise CensusWorkbookError("duplicate Census worksheet column labels")
                    label = f"{member.filename}:{sheet.title}"
                    return label, _subsection_rows(workbook, rows, headers, label, row_number)
        except Exception:
            workbook.close()
            raise
        workbook.close()
    raise CensusWorkbookError(
        "could not find a census synthesis table with required columns: "
        + ", ".join(sorted(required_columns))
        + " (XLSX also requires SUBSECCAO)"
    )


def _header(value: object) -> str:
    if not isinstance(value, str):
        return ""
    normalized = " ".join(unicodedata.normalize("NFKC", value).strip().split())
    return _HEADER_ALIASES.get(normalized, normalized)


def _text(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _subsection_rows(
    workbook: Workbook,
    rows: Iterator[Sequence[object]],
    headers: tuple[str, ...],
    label: str,
    header_row: int,
) -> Iterator[dict[str, str | None]]:
    parish_column = headers.index("DTMNFR21")
    subsection_column = headers.index("SUBSECCAO")
    seen: set[str] = set()
    try:
        for row_number, values in enumerate(rows, start=header_row + 1):
            subsection = _text(values[subsection_column])
            if not subsection:
                continue
            parish = _text(values[parish_column])
            if (
                parish is None
                or len(parish) != 6
                or not parish.isascii()
                or not parish.isalnum()
                or len(subsection) != 11
                or not subsection.startswith(parish)
                or not subsection[6:].isdigit()
            ):
                raise CensusWorkbookError(
                    f"{label}:row {row_number} has inconsistent parish/subsection identifiers"
                )
            if subsection in seen:
                raise CensusWorkbookError(f"{label}: duplicate subsection {subsection}")
            seen.add(subsection)
            yield {
                header: _text(value)
                for header, value in zip(headers, values, strict=True)
                if header
            }
    finally:
        workbook.close()
